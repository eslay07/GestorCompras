"""Read-only POP scanning, attachment review and non-overwriting local delivery."""
from __future__ import annotations

import email
import hashlib
import imaplib
import json
import os
import poplib
import re
import shutil
from datetime import datetime
from difflib import SequenceMatcher
from email.policy import default
from email.utils import getaddresses, parsedate_to_datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from sqlalchemy import text

from compras_tn.infrastructure.document_files import (
    DEFAULT_ROOT,
    file_types,
    inspect_file,
    matches_type,
    plain,
    task_folder,
)
from compras_tn.infrastructure.mail import decode_header_value

MONTHS = ("ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE")
LOCAL_ZONE = ZoneInfo("America/Guayaquil")


TELCONET_RUC = "0991327371001"


def local_document_text(path: Path) -> str:
    """Extract existing text only; never execute or modify local documents."""
    suffix = path.suffix.casefold()
    if suffix == '.pdf':
        import pdfplumber
        with pdfplumber.open(path) as document:
            return '\n'.join(page.extract_text() or '' for page in document.pages)
    if suffix in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return '\n'.join(' '.join(str(v) for v in row if v is not None)
                for sheet in workbook for row in sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    if suffix == '.docx':
        from zipfile import ZipFile
        from xml.etree import ElementTree
        with ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read('word/document.xml'))
        return ' '.join(root.itertext())
    if suffix in ('.txt', '.csv', '.tsv', '.xml', '.html', '.htm', '.eml'):
        content = path.read_bytes()
        try:
            return content.decode('utf-8-sig')
        except UnicodeDecodeError:
            return content.decode('cp1252', errors='replace')
    return ''


def document_rucs(text: str) -> set[str]:
    return {re.sub(r'\D', '', match.group()) for match in
            re.finditer(r'(?<!\d)\d(?:[ .-]?\d){12}(?!\d)', text)} - {TELCONET_RUC}


def safe_name(value: str, max_length: int = 120) -> str:
    result = re.sub(r'[<>:"/\\|?*\x00-\x1f]', " ", value)
    result = " ".join(result.split()).strip(" .")[:max_length].rstrip(" .")
    if not result:
        raise ValueError("No se pudo formar un nombre de archivo válido.")
    if result.split(".", 1)[0].upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        result = "_" + result
    return result


def quotation_subject(subject: str) -> bool:
    words = re.findall(r"[a-z]+", plain(subject))
    for index, word in enumerate(words):
        if SequenceMatcher(None, word, "solicitud").ratio() >= .85:
            following = words[index + 1:index + 4]
            if any(SequenceMatcher(None, candidate, "cotizacion").ratio() >= .85 for candidate in following):
                return bool(re.search(r"(?<!\d)\d{6,14}(?!\d)", subject))
    return False


class QuotationOrganizer:
    def __init__(self, database, settings, suppliers, sheets, paths, pop_factory=None):
        self.database, self.settings, self.suppliers, self.sheets = database, settings, suppliers, sheets
        self.cache = paths.root / "quotation-cache"
        self.pop_factory = pop_factory or poplib.POP3_SSL
        self.use_imap = pop_factory is None

    def root(self) -> Path:
        return Path(str(self.settings.get("quotations", "root", DEFAULT_ROOT))).resolve()

    def account(self, credentials) -> str:
        host = str(self.settings.get("downloads", "pop_server", "pop.telconet.ec"))
        return hashlib.sha256(f"{host.casefold()}|{credentials.email.casefold()}".encode()).hexdigest()

    def history(self, account: str) -> list[dict]:
        with self.database.session() as session:
            rows = session.execute(text("SELECT payload FROM quotation_attachments WHERE account=:account ORDER BY updated_at DESC"), {"account": account}).scalars()
            return [json.loads(row) for row in rows]

    def remember(self, row: dict) -> None:
        with self.database.session() as session:
            session.execute(text("""INSERT INTO quotation_attachments(identity,account,payload)
                VALUES (:identity,:account,:payload) ON CONFLICT(identity) DO UPDATE SET
                payload=excluded.payload,updated_at=CURRENT_TIMESTAMP"""),
                {"identity": row["identity"], "account": row["account"], "payload": json.dumps(row, ensure_ascii=False)})

    def month_dir(self, now: datetime | None = None) -> Path:
        now = (now or datetime.now(LOCAL_ZONE)).astimezone(LOCAL_ZONE)
        root = self.root()
        if not root.is_dir():
            raise FileNotFoundError(f"No existe la carpeta de destino: {root}")
        label = MONTHS[now.month - 1]
        candidates = [p for p in root.iterdir() if p.is_dir() and plain(p.name) == plain(label) and p.resolve().is_relative_to(root)]
        if len(candidates) > 1:
            raise ValueError("Hay varias carpetas para el mes actual. Revise la carpeta de destino.")
        return candidates[0] if candidates else root / label

    def destination(self, row: dict, now: datetime | None = None) -> tuple[Path | None, list[str]]:
        month = self.month_dir(now)
        if not row.get("project") or not re.fullmatch(r"[0-9]{6,14}", row.get("task_number", "")):
            raise ValueError("La tarea y su proyecto deben estar validados en GENERAL.")
        candidates = sorted([p.resolve() for p in month.iterdir() if p.is_dir() and task_folder(p.name, row["task_number"]) and p.resolve().is_relative_to(month.resolve())]) if month.is_dir() else []
        selected = row.get("chosen_folder", "")
        if selected:
            path = Path(selected).resolve()
            if path not in candidates:
                raise ValueError("La carpeta elegida no corresponde a esta tarea en el mes actual.")
            return path, []
        if len(candidates) > 1:
            return None, [str(p) for p in candidates]
        if candidates:
            return candidates[0], []
        budget = min(120, 220 - len(str(month)) - 65)
        if budget < len(row["task_number"]) + 5:
            raise ValueError("La ruta de destino es demasiado larga. Elija una carpeta raíz más corta.")
        return month / safe_name(row["task_number"] + " - " + row["project"], budget), []

    def validate_row(self, row: dict, catalog: dict) -> dict:
        number = row.get("task_number", "")
        info = catalog.get(number)
        row["project"] = info.get("project", "") if info else ""
        row["error"] = info.get("error", "") if info else "La tarea no existe en GENERAL."
        if row.get("ambiguous_subject"):
            row["error"] = "Seleccione la tarea correcta del reporte."
        if len(row.get("supplier_candidates", [])) > 1 and not row.get("supplier_confirmed"):
            row["error"] = ("El archivo contiene RUC de varios proveedores registrados. Revise el proveedor."
                            if row.get("source") == "local" else "Varios proveedores tienen este correo. Revise el proveedor.")
        row["folder_candidates"] = []
        row["planned_folder"] = ""
        if not row["error"]:
            try:
                path, choices = self.destination(row)
                row["folder_candidates"] = choices
                row["planned_folder"] = str(path) if path else ""
                if choices:
                    row["error"] = "Elija una carpeta de destino."
            except (OSError, ValueError) as exc:
                row["error"] = str(exc)
        target = Path(row.get("saved_path") or ".")
        saved = bool(row.get("saved_path") and target.is_file())
        if saved:
            try:
                saved = inspect_file(target)["sha256"] == row["sha256"]
            except (OSError, ValueError):
                saved = False
        row["status"] = "Pendiente de revisión" if row["error"] else "Ya guardado" if saved else "Restaurar archivo" if row.get("saved_path") else "Listo para guardar"
        return row

    def scan(self, credentials, start: datetime, end: datetime, tasks: list[str] | None = None, progress=None, cancelled=None) -> list[dict]:
        tasks = list(dict.fromkeys(tasks or []))
        if any(not re.fullmatch(r"[0-9]{6,14}", n) for n in tasks):
            raise ValueError("Ingrese números de tarea válidos, uno por línea.")
        catalog = self.sheets.read_task_catalog()
        account = self.account(credentials)
        previous = {row["identity"]: row for row in self.history(account)}
        found: dict[str, dict] = {}
        by_uid = {}
        for old_row in previous.values():
            by_uid.setdefault(old_row.get("uidl"), []).append(old_row)
        types = file_types(self.settings)
        suppliers = self.suppliers.list()
        supplier_emails = {}
        for supplier_record in suppliers:
            for _label, mail in getaddresses([supplier_record.email_primary.replace(";", ","), supplier_record.email_secondary.replace(";", ",")]):
                if mail.strip():
                    supplier_emails.setdefault(mail.strip().casefold(), []).append(supplier_record)
        if not supplier_emails:
            return []
        header_cache = self.settings.get("quotation_headers", account, {})
        if not isinstance(header_cache, dict):
            header_cache = {}
        host = str(self.settings.get("downloads", "pop_server", "pop.telconet.ec"))
        port = int(self.settings.get("downloads", "pop_port", 995))
        connection = None
        if self.use_imap and self.settings.get("quotations", "server_search", True):
            from compras_tn.infrastructure.quotation_mail import QuotationIMAP
            try:
                imap_host = str(self.settings.get("email", "imap_server", "pop.telconet.ec"))
                connection = QuotationIMAP(imap_host, int(self.settings.get("email", "imap_port", 993)), list(supplier_emails), start, end, tasks)
                connection.user(credentials.email)
                connection.pass_(credentials.password)
            except (OSError, imaplib.IMAP4.error):
                if connection:
                    try:
                        connection.quit()
                    except (OSError, imaplib.IMAP4.error):
                        pass
                connection = None
                if progress:
                    progress(0, 0, "Busqueda rapida no disponible; usando POP y cache local.")
        if connection is None:
            connection = self.pop_factory(host, port, timeout=30)
            connection.user(credentials.email)
            connection.pass_(credentials.password)
        try:
            _response, lines, _size = connection.uidl()
            for index, line in enumerate(reversed(lines), 1):
                if cancelled and cancelled():
                    break
                if progress:
                    progress(index, len(lines), f"Revisando correo {index} de {len(lines)}")
                parts = line.decode(errors="replace").split(maxsplit=1)
                if len(parts) != 2:
                    continue
                message_number, uid = parts
                cached_header = header_cache.get(uid)
                if cached_header:
                    headers = cached_header
                else:
                    try:
                        _, header_lines, _ = connection.top(int(message_number), 0)
                    except (poplib.error_proto, AttributeError):
                        _, header_lines, _ = connection.retr(int(message_number))
                    parsed_headers = email.message_from_bytes(b"\r\n".join(header_lines), policy=default)
                    headers = {key: str(parsed_headers.get(key, "")) for key in ("Subject", "From", "Date")}
                    header_cache[uid] = headers
                addresses = getaddresses([decode_header_value(headers.get("From", ""))])
                display, address = addresses[0] if addresses else ("", "")
                candidates = supplier_emails.get(address.strip().casefold(), [])
                if not candidates:
                    continue
                subject = decode_header_value(headers.get("Subject", ""))
                numbers = list(dict.fromkeys(re.findall(r"(?<!\d)\d{6,14}(?!\d)", subject)))
                if tasks and not set(numbers).intersection(tasks):
                    continue
                if not quotation_subject(subject):
                    continue
                try:
                    received = parsedate_to_datetime(str(headers.get("Date", "")))
                    if received.tzinfo is None:
                        received = received.replace(tzinfo=LOCAL_ZONE)
                except (TypeError, ValueError, OverflowError):
                    received = None
                if not tasks and (not received or not start <= received <= end):
                    continue
                eligible = [n for n in numbers if n in catalog]
                if not eligible:
                    continue
                matched = eligible
                old_rows = by_uid.get(uid, [])
                if old_rows and all(Path(row.get("cached_path", "")).is_file() for row in old_rows):
                    for old_row in old_rows:
                        if tasks and old_row.get("task_number") not in tasks and not set(old_row.get("task_candidates", [])).intersection(tasks):
                            continue
                        row = dict(old_row)
                        if row.get("task_candidates") != matched:
                            row.update(task_candidates=matched, task_number=matched[0] if len(matched) == 1 else "", ambiguous_subject=len(matched) > 1)
                        if len(candidates) == 1 and not row.get("supplier_confirmed"):
                            row["supplier"] = candidates[0].name
                        row = self.validate_row(row, catalog)
                        found[row["identity"]] = row
                    continue
                _, full_lines, _ = connection.retr(int(message_number))
                message = email.message_from_bytes(b"\r\n".join(full_lines), policy=default)
                number = matched[0] if len(matched) == 1 else ""
                supplier = candidates[0].name if len(candidates) == 1 else display.strip() or address.split("@", 1)[0] or "Remitente desconocido"
                for part_index, part in enumerate(message.walk()):
                    filename = decode_header_value(part.get_filename() or "")
                    if not filename or part.is_multipart() or (part.get_content_disposition() != "attachment" and part.get_content_type().startswith("image/")):
                        continue
                    content = part.get_payload(decode=True)
                    if not content:
                        continue
                    digest = hashlib.sha256(content).hexdigest()
                    identity = hashlib.sha256(f"{account}|{uid}|{part_index}".encode()).hexdigest()
                    self.cache.mkdir(parents=True, exist_ok=True)
                    cached = self.cache / digest
                    if not cached.is_file() or inspect_file(cached)["sha256"] != digest:
                        temp = self.cache / (identity + ".tmp")
                        with temp.open("wb") as stream:
                            stream.write(content)
                            stream.flush()
                            os.fsync(stream.fileno())
                        temp.replace(cached)
                    recognized = [r["id"] for r in types if matches_type(Path(filename), r)]
                    row = {"identity": identity, "account": account, "uidl": uid,
                           "message_id": str(message.get("Message-ID", "")), "subject": subject,
                           "task_number": number, "task_candidates": matched, "ambiguous_subject": len(matched) > 1,
                           "sender": address, "supplier": supplier, "supplier_candidates": [p.name for p in candidates], "received": received.astimezone(LOCAL_ZONE).isoformat() if received else "",
                           "original_name": filename, "type_id": recognized[0] if len(recognized) == 1 else "otros",
                           "cached_path": str(cached), "sha256": digest, "size": len(content), "saved_path": ""}
                    old = previous.get(identity)
                    if old and old.get("sha256") == digest:
                        for key in ("supplier", "type_id", "saved_path", "chosen_folder", "supplier_confirmed"):
                            if old.get(key):
                                row[key] = old[key]
                    if old and old.get("task_number") in matched and not old.get("ambiguous_subject"):
                        row["task_number"] = old["task_number"]
                        row["ambiguous_subject"] = False
                    row = self.validate_row(row, catalog)
                    self.remember(row)
                    found[identity] = row
            if tasks:
                for row in previous.values():
                    if row["identity"] not in found and row.get("task_number") in tasks and row.get("task_number") in catalog and row.get("sender", "").casefold() in supplier_emails and quotation_subject(row.get("subject", "")):
                        row = self.validate_row(row, catalog)
                        if not Path(row.get("cached_path", "")).is_file():
                            row["error"], row["status"] = "El adjunto no está disponible en caché ni en el servidor.", "Pendiente de revisión"
                        found[row["identity"]] = row
            return list(found.values())
        finally:
            self.settings.set("quotation_headers", account, header_cache)
            try:
                connection.quit()
            except (OSError, poplib.error_proto, imaplib.IMAP4.error):
                pass

    def review_local_tasks(self, rows, progress=None, cancelled=None):
        catalog = self.sheets.read_task_catalog()
        scanned = self.scan_local([row['cached_path'] for row in rows if row.get('cached_path')],
                                  progress, cancelled, catalog=catalog)
        by_path = {row.get('cached_path'): row for row in scanned if row.get('cached_path')}
        result = []
        for incoming in rows:
            if cancelled and cancelled():
                break
            fresh = by_path.get(incoming.get('cached_path'))
            row = dict(incoming)
            if fresh:
                for key in ('supplier', 'supplier_candidates', 'supplier_ruc', 'detected_rucs', 'read_warning'):
                    if not incoming.get('supplier_confirmed'):
                        row[key] = fresh.get(key)
            row.pop('chosen_folder', None)
            row.pop('planned_path', None)
            row.pop('planned_name', None)
            row['ambiguous_subject'] = False
            row = self.validate_row(row, catalog)
            if not fresh or fresh.get('error') and not fresh.get('cached_path'):
                row['error'] = 'El archivo local ya no está disponible.'
            if not row.get('supplier') and not row['error']:
                row['error'] = 'No se identificó un proveedor registrado. Revise el proveedor.'
            result.append(row)
        return result

    def review_local_task(self, row: dict) -> dict:
        updated = dict(row)
        updated.pop('chosen_folder', None)
        updated.pop('planned_path', None)
        updated.pop('planned_name', None)
        updated['ambiguous_subject'] = False
        return self.validate_row(updated, self.sheets.read_task_catalog())

    def scan_local(self, paths, progress=None, cancelled=None, *, catalog=None) -> list[dict]:
        """Review local originals through the same validation and delivery pipeline."""
        catalog = self.sheets.read_task_catalog() if catalog is None else catalog
        types = file_types(self.settings)
        suppliers = self.suppliers.list()
        result = []
        paths = list(dict.fromkeys(str(Path(p).resolve()) for p in paths))
        for index, raw in enumerate(paths, 1):
            if cancelled and cancelled():
                break
            path = Path(raw)
            try:
                info = inspect_file(path)
                numbers = re.findall(r"(?<!\d)\d{6,14}(?!\d)", path.stem)
                matched = list(dict.fromkeys(n for n in numbers if n in catalog))
                if not matched:
                    matched = [n for n in catalog if task_folder(path.parent.name, n)]
                read_warning = ''
                try:
                    content = local_document_text(path)
                except Exception:
                    content = ''
                    read_warning = 'No se pudo leer el contenido; revise el proveedor manualmente.'
                rucs = document_rucs(content)
                filename_rucs = document_rucs(path.stem)
                # Content has priority over the filename and customer data.
                lookup = rucs or filename_rucs
                eligible = [p for p in suppliers if re.sub(r'\D', '', p.ruc) != TELCONET_RUC]
                if lookup:
                    matches = [p for p in eligible if re.sub(r'\D', '', p.ruc) in lookup]
                elif TELCONET_RUC in re.sub(r'\D', '', content):
                    matches = []
                else:
                    matches = [p for p in eligible if plain(p.name) in plain(path.stem)]
                candidates = [p.name for p in matches]
                recognized = [t['id'] for t in types if matches_type(path, t)]
                identity = hashlib.sha256(f"local|{raw}|{info['sha256']}".encode()).hexdigest()
                row = {
                    'identity': identity, 'account': 'local', 'source': 'local',
                    'task_number': matched[0] if len(matched) == 1 else '',
                    'task_candidates': matched or sorted(catalog),
                    'ambiguous_subject': len(matched) != 1,
                    'supplier': candidates[0] if len(candidates) == 1 else '',
                    'supplier_candidates': candidates,
                    'supplier_ruc': matches[0].ruc if len(matches) == 1 else '',
                    'detected_rucs': sorted(lookup), 'read_warning': read_warning,
                    'received': datetime.now(LOCAL_ZONE).isoformat(),
                    'original_name': path.name,
                    'type_id': recognized[0] if len(recognized) == 1 else 'otros',
                    'cached_path': raw, 'sha256': info['sha256'],
                    'size': info['size'], 'saved_path': '',
                }
                row = self.validate_row(row, catalog)
                if not row['supplier'] and not row['error']:
                    row['status'] = read_warning or ('RUC no registrado: ' + ', '.join(sorted(lookup)) if lookup else 'Complete el proveedor antes de guardar')
                result.append(row)
            except (OSError, ValueError) as exc:
                result.append({'identity': hashlib.sha256(raw.encode()).hexdigest(),
                    'account': 'local', 'original_name': path.name, 'type_id': 'otros',
                    'task_number': '', 'supplier': '', 'error': str(exc),
                    'status': 'No disponible'})
            if progress:
                progress(index, len(paths), f"Revisando archivo local {index} de {len(paths)}")
        return result

    def planned_target(self, row: dict, folder: Path, types: dict, reserved: dict) -> Path:
        if folder.is_dir():
            for existing in folder.iterdir():
                if existing.is_file() and not existing.name.endswith(".tmp"):
                    try:
                        if inspect_file(existing)["sha256"] == row["sha256"]:
                            return existing
                    except (OSError, ValueError):
                        continue
        for path, digest in reserved.items():
            if Path(path).parent == folder and digest == row["sha256"]:
                return Path(path)
        kind = plain(types[row["type_id"]]["name"]).upper()
        extension = Path(row["original_name"]).suffix
        if extension and (len(extension) > 30 or any(c in extension for c in '<>:"/\\|?*')):
            raise ValueError("La extensión original no es válida para Windows.")
        budget = min(100, 238 - len(str(folder)) - len(extension) - 1)
        if budget < len(kind) + 5:
            raise ValueError("La ruta es demasiado larga para conservar el tipo del documento.")
        prefix = safe_name(kind + " - " + " ".join(row["supplier"].split()[:2]).upper(), budget)
        target = folder / (prefix + extension)
        stamp = row.setdefault("version_stamp", datetime.now(LOCAL_ZONE).strftime("%Y-%m-%d %H%M%S"))
        serial = 1
        while target.exists() or str(target) in reserved:
            shortened = prefix[:max(len(kind) + 5, len(prefix) - 24)]
            target = folder / (shortened + f" - {stamp} - {serial:02d}" + extension)
            if len(str(target)) > 250:
                raise ValueError("La ruta de la versión es demasiado larga. Elija una raíz más corta.")
            serial += 1
        reserved[str(target)] = row["sha256"]
        return target

    def prepare_save(self, rows: list[dict], *, confirmed: bool = False) -> list[dict]:
        catalog = self.sheets.read_task_catalog()
        types = {r["id"]: r for r in file_types(self.settings)}
        reserved: dict[str, str] = {}
        result = []
        for incoming in rows:
            row = dict(incoming)
            reviewed_folder = row.get("planned_folder", "")
            row = self.validate_row(row, catalog)
            if reviewed_folder and row["planned_folder"] != reviewed_folder:
                raise ValueError("El destino cambió desde la revisión. Actualice el reporte antes de guardar.")
            if row["error"]:
                raise ValueError(f"Tarea {row.get('task_number') or '?'}: {row['error']}")
            source = inspect_file(Path(row["cached_path"]))
            if source["sha256"] != row["sha256"]:
                raise ValueError("El adjunto cambió desde su revisión. Escanee nuevamente.")
            if not row.get("supplier", "").strip():
                raise ValueError("Complete el proveedor o remitente.")
            if row["type_id"] not in {r["id"] for r in file_types(self.settings)}:
                raise ValueError("Seleccione un tipo de documento válido.")
            target = self.planned_target(row, Path(row["planned_folder"]), types, reserved)
            if confirmed and incoming.get("planned_path") and str(target) != incoming["planned_path"]:
                raise ValueError("El nombre de destino cambió desde la confirmación. Actualice el reporte.")
            row["planned_name"], row["planned_path"] = target.name, str(target)
            result.append(row)
        return result

    def save(self, rows: list[dict], progress=None, cancelled=None) -> list[dict]:
        rows = self.prepare_save(rows, confirmed=True)
        result = []
        for index, row in enumerate(rows, 1):
            if cancelled and cancelled():
                break
            try:
                # Re-read the cache immediately before copying.
                source = inspect_file(Path(row["cached_path"]))
                if source["sha256"] != row["sha256"]:
                    raise ValueError("El adjunto cambió antes de guardarse.")
                folder = Path(row["planned_folder"])
                if not folder.resolve().is_relative_to(self.root()):
                    raise ValueError("El destino salió de la carpeta general.")
                folder.mkdir(parents=True, exist_ok=True)
                if not folder.resolve().is_relative_to(self.root()):
                    raise ValueError("El destino salió de la carpeta general.")
                target = Path(row["planned_path"])
                if target.is_file() and inspect_file(target)["sha256"] == row["sha256"]:
                    row.update(saved_path=str(target), status="Ya guardado")
                else:
                    row.update(status="Guardando", planned_path=str(target))
                    self.remember(row)
                    temp = folder / (".quotation-" + row["identity"][:12] + ".tmp")
                    try:
                        with Path(row["cached_path"]).open("rb") as src, temp.open("wb") as dest:
                            shutil.copyfileobj(src, dest)
                            dest.flush()
                            os.fsync(dest.fileno())
                        if inspect_file_for_temp(temp) != row["sha256"]:
                            raise ValueError("La copia no coincide con el adjunto original.")
                        # Publish atomically without replacing existing content.
                        if os.name == "nt":
                            os.rename(temp, target)
                        else:
                            os.link(temp, target)
                        row.update(saved_path=str(target), status="Guardado")
                    finally:
                        temp.unlink(missing_ok=True)
                row["error"] = ""
                self.remember(row)
            except (OSError, ValueError) as exc:
                row.update(error=str(exc), status="Error al guardar")
                self.remember(row)
            result.append(row)
            if progress:
                progress(index, len(rows), f"{row['task_number']}: {row['status']}")
        return result


def inspect_file_for_temp(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()
