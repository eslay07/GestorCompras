from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from compras_tn.domain.models import OperationRequest, OperationResult
from compras_tn.paths import AppPaths
from compras_tn.security import sanitize_text

ILLEGAL_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")


class ExcelReportGateway:
    def __init__(self, paths: AppPaths):
        self.paths = paths

    def write(self, request: OperationRequest, result: OperationResult) -> Path:
        self.paths.reports.mkdir(parents=True, exist_ok=True)
        safe_kind = re.sub(r"[^A-Za-z0-9_-]+", "_", request.kind).strip("_") or "operacion"
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        destination = self.paths.reports / f"{safe_kind}_{stamp}.xlsx"
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Resultados"
        sheet.append(["Operación", self._cell(request.kind)])
        sheet.append(["Modo", self._cell(request.mode.value)])
        sheet.append(["Estado", self._cell(result.status.value)])
        sheet.append(["Inicio", result.started_at.isoformat()])
        sheet.append(["Fin", result.finished_at.isoformat() if result.finished_at else ""])
        sheet.append(["Resumen", self._cell(result.message)])
        sheet.append([])
        purchase_orders = request.kind == "precarga_oc_bienes"
        headers = (
            [
                "Elemento",
                "Tarea",
                "Proveedor",
                "RUC",
                "Solicitante NAF",
                "OC asignada",
                "Artículos",
                "Subtotal",
                "Estado",
                "Código",
                "Detalle",
            ]
            if purchase_orders
            else ["Elemento", "Estado", "Código", "Detalle"]
        )
        sheet.append(headers)
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
        for item in result.items:
            if purchase_orders:
                sheet.append([
                    self._cell(item.key),
                    self._cell(item.payload.get("task_number", "")),
                    self._cell(item.payload.get("supplier", "")),
                    self._cell(item.payload.get("ruc", "")),
                    self._cell(item.payload.get("requester", "")),
                    self._cell(item.payload.get("order_number", "")),
                    self._cell(item.payload.get("articles", "")),
                    self._cell(item.payload.get("subtotal", "")),
                    self._cell(item.status.value),
                    self._cell(item.error_code),
                    self._cell(item.message),
                ])
            else:
                sheet.append([
                    self._cell(item.key), self._cell(item.status.value),
                    self._cell(item.error_code), self._cell(item.message),
                ])
        for column, width in (
            ("A", 28), ("B", 18), ("C", 30), ("D", 18), ("E", 32),
            ("F", 18), ("G", 12), ("H", 14), ("I", 18), ("J", 28), ("K", 90),
        ):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = f"A{header_row + 1}"
        temporary = destination.with_suffix(".tmp.xlsx")
        try:
            workbook.save(temporary)
            temporary.replace(destination)
        finally:
            workbook.close()
            if temporary.exists():
                temporary.unlink()
        return destination

    @staticmethod
    def _cell(value: object) -> str:
        text = ILLEGAL_EXCEL_CHARACTERS.sub("", sanitize_text(value))[:32_767]
        if text.startswith(("=", "+", "-", "@")):
            return f"'{text}"
        return text
